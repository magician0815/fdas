"""
NY Fed Excel 宏观数据采集器.

处理 r-star-LW (Laubach-Williams) 和 r-star-HLW (Holston-Laubach-Williams) 两类自然利率数据.
数据源: https://www.newyorkfed.org/research/policy/rstar

采集流程: 先获取研究页面 -> 解析页面提取 Excel 下载链接 -> 下载 Excel 文件 -> 解析

Author: FDAS Team
Created: 2026-07-29
Updated: 2026-07-29 - 两阶段下载: HTML 页面 → Excel 文件
"""

import logging
import re
from io import BytesIO
from typing import Any, Dict, Optional

import httpx
import pandas as pd
import openpyxl
from bs4 import BeautifulSoup

from app.collectors.macro_base_collector import MacroBaseCollector

logger = logging.getLogger(__name__)


class MacroNYFedExcelCollector(MacroBaseCollector):
    """
    NY Fed Excel 数据采集器.

    支持 LW 模型(美国单国)和 HLW 模型(多国)数据解析.
    采集分两步: 先获取 rstar 研究首页, 从中提取 Excel 文件下载链接.

    parse_config 关键字段:
        sheet_name: Excel 工作表名
        skiprows: 跳过行数
        columns: 源列名 -> 目标列名映射
        indicator_key: 指标键名
        country: 默认国家
        multi_country: 是否为多国数据
        download_pattern: Excel 文件 URL 匹配模式(如 "lw_data.xlsx")
    """

    async def _fetch_raw(self) -> bytes:
        """重写: 两步获取实际 Excel 数据."""
        # 第一步: 获取研究页面, 提取 Excel 下载链接
        async with httpx.AsyncClient(timeout=self.timeout) as client:
            resp = await client.get(self.url, headers=self.headers)
            resp.raise_for_status()

        soup = BeautifulSoup(resp.content, "lxml")
        download_pattern = self.parse_config.get("download_pattern", ".xlsx")
        excel_url = None

        # 查找匹配的 Excel 下载链接
        for a_tag in soup.find_all("a", href=True):
            href = a_tag["href"]
            if download_pattern.lower() in href.lower():
                if href.startswith("http"):
                    excel_url = href
                else:
                    excel_url = f"https://www.newyorkfed.org{href}"
                break

        if not excel_url:
            raise ValueError(f"未找到匹配 '{download_pattern}' 的 Excel 下载链接")

        logger.info(f"[{self.source_code}] Excel 下载链接: {excel_url}")

        # 第二步: 下载 Excel 文件
        async with httpx.AsyncClient(timeout=self.timeout) as client:
            resp = await client.get(excel_url, headers=self.headers)
            resp.raise_for_status()
            return resp.content

    async def parse_raw(self, raw_data: bytes) -> pd.DataFrame:
        wb = openpyxl.load_workbook(BytesIO(raw_data), data_only=True)
        sheet_name = self.parse_config.get("sheet_name") or wb.sheetnames[0]
        ws = wb[sheet_name]

        skiprows = self.parse_config.get("skiprows", 0)
        col_mapping = self.parse_config.get("columns", {})
        frequency = self.parse_config.get("frequency", "quarterly")
        indicator_key = self.parse_config.get("indicator_key", self.source_code)
        multi_country = self.parse_config.get("multi_country", False)
        default_country = self.parse_config.get("country", "US")

        country_columns = self.parse_config.get("country_columns", {})

        # 模式1: 按列索引映射(HLW 多段表格, 处理重复列名)
        if multi_country and country_columns:
            records = self._parse_index_based(ws, skiprows, col_mapping, indicator_key, frequency, country_columns)

        # 模式2: 按列名映射(LW 单国 或 含 country 列)
        else:
            records = self._parse_header_based(ws, skiprows, col_mapping, indicator_key, frequency, multi_country, default_country)

        logger.info(f"[{self.source_code}] Excel 解析完成: {len(records)} 条")
        return pd.DataFrame(records)

    def _parse_index_based(self, ws, skiprows, col_mapping, indicator_key, frequency, country_columns):
        """按列索引解析(处理多段表头和重复列名)."""
        date_col_idx = col_mapping.get("date_col", 1)

        records = []
        for row_idx in range(skiprows + 2, ws.max_row + 1):
            date_val = ws.cell(row=row_idx, column=date_col_idx).value
            if date_val is None:
                continue

            for country_code, col_idx in country_columns.items():
                actual_value = ws.cell(row=row_idx, column=col_idx).value
                if actual_value is None:
                    continue
                if isinstance(actual_value, str) and actual_value.strip().upper() in ("NA", "N/A", ""):
                    continue
                try:
                    float(actual_value)
                except (TypeError, ValueError):
                    continue

                records.append({
                    "country": country_code,
                    "series_name": f"{indicator_key.upper()} {country_code}",
                    "indicator_key": indicator_key,
                    "value": actual_value,
                    "publish_date": None,
                    "period_date": date_val,
                    "frequency": frequency,
                })

        return records

    def _parse_header_based(self, ws, skiprows, col_mapping, indicator_key, frequency, multi_country, default_country):
        """按列名解析(LW 单国)."""
        headers = []
        for col_idx, cell in enumerate(ws[skiprows + 1], start=1):
            headers.append(cell.value)

        rows = []
        for row_idx in range(skiprows + 2, ws.max_row + 1):
            row_data = {}
            for col_idx, cell in enumerate(ws[row_idx], start=1):
                if col_idx <= len(headers):
                    row_data[headers[col_idx - 1]] = cell.value
            rows.append(row_data)

        df = pd.DataFrame(rows)
        if df.empty:
            return []

        date_col = col_mapping.get("date")
        value_col = col_mapping.get("r_star_hlw") or col_mapping.get("r_star_lw")
        country_col = col_mapping.get("country")

        if date_col not in df.columns:
            raise ValueError(f"日期列 '{date_col}' 不在 Excel 中, 可用列: {list(df.columns)}")
        if value_col and value_col not in df.columns:
            raise ValueError(f"数值列 '{value_col}' 不在 Excel 中, 可用列: {list(df.columns)}")

        records = []
        for _, row in df.iterrows():
            date_val = row.get(date_col)
            if date_val is None:
                continue

            countries_to_process = [default_country]
            if multi_country and country_col and country_col in df.columns:
                c = row.get(country_col)
                if c:
                    countries_to_process = [str(c)]

            for country in countries_to_process:
                actual_value = row.get(value_col)
                if actual_value is None:
                    continue
                records.append({
                    "country": country,
                    "series_name": f"{indicator_key.upper()} {country}",
                    "indicator_key": indicator_key,
                    "value": actual_value,
                    "publish_date": None,
                    "period_date": date_val,
                    "frequency": frequency,
                })

        return records

        logger.info(f"[{self.source_code}] Excel 解析完成: {len(records)} 条")
        return pd.DataFrame(records)
